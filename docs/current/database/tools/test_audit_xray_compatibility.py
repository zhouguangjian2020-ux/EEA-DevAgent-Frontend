"""Unit checks for read-only compatibility extraction. No database or network."""
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl
from audit_xray_compatibility import NS, W, text, read_doc, cell_value_with_explicit_merge, SUGGESTIONS


class CompatibilityExtractionTests(unittest.TestCase):
    def test_explicit_merge_is_expanded(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['B3'] = '车身\n（1XX）'
        sheet.merge_cells('B3:B5')
        self.assertEqual(cell_value_with_explicit_merge(sheet, 5, 2), '车身\n（1XX）')
        book.close()

    def test_unmerged_blank_is_not_forward_filled(self):
        book = openpyxl.Workbook()
        book.active['B3'] = '车身'
        self.assertIsNone(cell_value_with_explicit_merge(book.active, 4, 2))
        book.close()

    def test_original_whitespace_symbols_and_breaks_preserved(self):
        node = ET.fromstring(f'<w:p xmlns:w="{NS["w"]}"><w:r><w:t xml:space="preserve"> N/A ≤500ms（TBD） </w:t><w:br/><w:t>无</w:t><w:tab/><w:t>末</w:t></w:r></w:p>')
        self.assertEqual(text(node), ' N/A ≤500ms（TBD） \n无\t末')

    def parse_fixture(self, table):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'fixture.docx'
            with ZipFile(path, 'w') as package:
                package.writestr('word/document.xml', f'<w:document xmlns:w="{NS["w"]}"><w:body>{table}</w:body></w:document>')
            return read_doc(path)

    def test_word_horizontal_and_vertical_merges(self):
        doc = self.parse_fixture('<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>场景</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr></w:tbl>')
        self.assertEqual(doc['tables'], [[['场景', '场景'], ['场景', '场景']]])

    def test_invalid_vertical_merge_fails_explicitly(self):
        with self.assertRaisesRegex(ValueError, 'Unmatched vertical merge'):
            self.parse_fixture('<w:tbl><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr></w:tbl>')

    def test_automatic_numbering_is_reported_not_invented(self):
        doc = self.parse_fixture('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>步骤</w:t></w:r></w:p>')
        self.assertEqual(doc['explicit_numbering_properties'], 1)
        self.assertEqual(doc['plain_text'], '步骤')

    def test_review_suggestions_are_explicit_not_fuzzy(self):
        self.assertEqual(len(SUGGESTIONS), 5)
        self.assertEqual(SUGGESTIONS['X-Call系统'], 'X-call系统')
        self.assertNotIn('E-Call系统', SUGGESTIONS)
        self.assertNotIn('其他系统', SUGGESTIONS)


if __name__ == '__main__':
    unittest.main(verbosity=2)
