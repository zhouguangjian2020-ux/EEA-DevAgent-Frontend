"""Read-only XLSX/DOCX compatibility audit; writes local evidence, never SQL/network.

Run with the bundled Python runtime. Source files and old DVR staging are unchanged.
Candidate mappings are suggestions for human review, NOT automatic alias rules.
"""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from collections import Counter
import argparse
import hashlib
import json
import re
import openpyxl

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'
BASE = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r'C:\Users\zhouguangjian\Desktop\X光\数据')
# Explicit review suggestions for this supplied sample, never fuzzy-match or auto-apply.
SUGGESTIONS = {
    '低压电源子系统': '低压电源管理系统',
    '感知融合与定位系统': '感知融合和定位系统',
    '挡位系统': '挡位控制系统',
    'X-Call系统': 'X-call系统',
    '外部连接子系统': '外部设备连接系统',
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(element):
    return ''.join(n.text or '' if n.tag == W + 't' else
                   '\n' if n.tag in (W + 'br', W + 'cr') else '\t'
                   for n in element.iter() if n.tag in
                   (W + 't', W + 'br', W + 'cr', W + 'tab'))


def read_doc(path):
    with ZipFile(path) as package:
        root = ET.fromstring(package.read('word/document.xml'))
        body = root.find('w:body', NS)
        tables = []
        for table in body.findall('w:tbl', NS):
            rows, anchors = [], {}
            for tr in table.findall('w:tr', NS):
                row, col = [], 0
                for tc in tr.findall('w:tc', NS):
                    span = tc.find('w:tcPr/w:gridSpan', NS)
                    span = int(span.get(W + 'val')) if span is not None else 1
                    merge = tc.find('w:tcPr/w:vMerge', NS)
                    mode = merge.get(W + 'val', 'continue') if merge is not None else None
                    value = '\n'.join(text(p) for p in tc.findall('w:p', NS))
                    if mode == 'continue':
                        if col not in anchors:
                            raise ValueError('Unmatched vertical merge')
                        value = anchors[col]
                    for offset in range(span):
                        row.append(value)
                        if mode in ('restart', 'continue'):
                            anchors[col + offset] = value
                        else:
                            anchors.pop(col + offset, None)
                    col += span
                rows.append(row)
            tables.append(rows)
        return {
            'tables': tables,
            'nonempty_body_paragraphs': sum(bool(text(p)) for p in body.findall('w:p', NS)),
            'explicit_numbering_properties': len(body.findall('.//w:numPr', NS)),
            'tracked_insertions': len(body.findall('.//w:ins', NS)),
            'tracked_deletions': len(body.findall('.//w:del', NS)),
            'body_drawings': len(body.findall('.//w:drawing', NS)),
            'body_pictures': len(body.findall('.//w:pict', NS)),
            'package_media': [n for n in package.namelist() if n.startswith('word/media/')],
            'plain_text': text(body),
        }


def cell_value_with_explicit_merge(sheet, row, col):
    value = sheet.cell(row, col).value
    if value is not None:
        return value
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return None  # Never forward-fill an unmerged blank.


def audit(source, output):
    required = ['PRD文档.docx', 'UC.docx', 'FR列表.xlsx', 'FR划分.xlsx', '系统划分.xlsx']
    before = {name: digest(source / name) for name in required}
    old = json.loads((BASE / 'seed-candidates/dvr-20260921/audit.json').read_text(encoding='utf-8'))
    prior = {s['name']: s['sha256'] for s in old['sources']}
    workbooks = {name: openpyxl.load_workbook(source / name, data_only=False)
                 for name in required if name.endswith('.xlsx')}
    try:
        ws = workbooks['系统划分.xlsx']['系统功能需求']
        if [ws.cell(2, col).value for col in (2, 3, 4)] != ['域', '系统名称', '系统描述']:
            raise ValueError('Unexpected subsystem header')
        items = []
        for row in range(3, ws.max_row + 1):
            name = ws.cell(row, 3).value
            if name is None:
                continue
            items.append({
                'subsystem_id': None, 'code': None,
                'name': name, 'description': ws.cell(row, 4).value,
                'domain_group': cell_value_with_explicit_merge(ws, row, 2),
                'responsibility_in': None, 'responsibility_out': None,
                'sort_order': len(items), 'is_allocatable': None,
                'source_metadata': {'file': '系统划分.xlsx', 'sha256': before['系统划分.xlsx'],
                                    'sheet': ws.title, 'range': f'B{row}:D{row}'}
            })
        def records(file, sheet, columns):
            s = workbooks[file][sheet]
            return [(row, tuple(s.cell(row, col).value for col in range(1, columns + 1)))
                    for row in range(2, s.max_row + 1) if s.cell(row, 1).value is not None]
        fr_a = records('FR列表.xlsx', '行车记录仪Fr_List', 3)
        fr_b = records('FR划分.xlsx', '行车记录仪Fr_List', 4)
        alloc_a = records('FR列表.xlsx', '行车记录仪功能需求分配表 ', 3)
        alloc_b = records('FR划分.xlsx', '行车记录仪功能需求分配表 ', 3)
        uc_doc = read_doc(source / 'UC.docx')
        prd_doc = read_doc(source / 'PRD文档.docx')
        ucs = []
        for index, table in enumerate(uc_doc['tables'], 1):
            fields, branches, exceptions = {}, [], []
            for row in table:
                label = row[0].strip().rstrip('：:')
                if label == '场景':
                    if row[1].startswith('分支事件流'):
                        branches.append(row[-1])
                    elif row[1].startswith('异常事件流'):
                        exceptions.append({'category': row[2], 'body': row[-1]})
                else:
                    fields[label] = row[-1]
            ucs.append({'table': index, 'code': fields['用例编号'],
                        'name': fields['用例名称'], 'top_level_labels': list(fields) + ['场景'],
                        'branch_count': len(branches), 'exception_count': len(exceptions),
                        'blank_hmi_raw': fields['HMI 交互逻辑'],
                        'text_uc_mentions': sorted(set(re.findall(r'UC-\d+-\d+-\d+', '\n'.join(r[-1] for r in table))))})
        uc_codes = {u['code'] for u in ucs}
        fr_sources = [{'fr_code': values[0], 'uc_codes': re.findall(r'UC-\d+-\d+-\d+', values[1])}
                      for _, values in fr_a]
        by_name = {item['name']: item for item in items}
        allocations = Counter(values[1] for _, values in alloc_b)
        mapping = []
        for name, count in allocations.items():
            exact = name in by_name
            suggestion = name if exact else SUGGESTIONS.get(name)
            candidate = by_name.get(suggestion)
            mapping.append({'source_name': name, 'target_count': count,
                            'fr_codes': [v[2] for _, v in alloc_b if v[1] == name],
                            'status': 'exact_name_candidate' if exact else 'needs_human_confirmation',
                            'suggested_catalog_name': suggestion,
                            'catalog_range': candidate['source_metadata']['range'] if candidate else None,
                            'confirmed': False})
        checks = {
            'fr_copies_equal': [v for _, v in fr_a] == [v[:3] for _, v in fr_b],
            'allocation_copies_equal': alloc_a == alloc_b,
            'allocation_D_matches_table': {v[0]: v[3] for _, v in fr_b} == {v[2]: v[1] for _, v in alloc_b},
            'unique_fr_codes': len({v[0] for _, v in fr_a}) == len(fr_a),
            'unique_uc_codes': len(uc_codes) == len(ucs),
            'all_fr_uc_sources_resolve': all(s in uc_codes for row in fr_sources for s in row['uc_codes']),
            'allocation_fr_set_matches': {v[2] for _, v in alloc_b} == {v[0] for _, v in fr_a},
            'catalog_names_unique': len(by_name) == len(items),
            'catalog_fields_present': all(i['name'] and i['description'] and i['domain_group'] for i in items),
            'no_formula_cells': not any(c.data_type == 'f' for book in workbooks.values() for s in book for row in s for c in row),
            'old_four_source_hashes_unchanged': all(before[n] == prior.get(n) for n in required if n != '系统划分.xlsx'),
        }
        unresolved = sorted({s for u in ucs for s in u['text_uc_mentions']} - uc_codes)
        summary = {
            'status': 'audit_only_not_imported', 'inspected_on': '2026-09-21',
            'source_manifest': [{'file': str(source / name), 'sha256': before[name],
                                 'previously_inventoried': name in prior} for name in required],
            'counts': {'prd_documents': 1, 'uc': len(ucs), 'fr': len(fr_a),
                       'fr_uc_links': sum(len(r['uc_codes']) for r in fr_sources),
                       'allocation_targets': len(alloc_b), 'allocation_names': len(allocations),
                       'catalog_items': len(items), 'domain_groups': len({i['domain_group'] for i in items}),
                       'exact_name_targets': sum(m['target_count'] for m in mapping if m['status'] == 'exact_name_candidate'),
                       'unresolved_name_targets': sum(m['target_count'] for m in mapping if m['status'] != 'exact_name_candidate')},
            'domain_counts': dict(Counter(i['domain_group'] for i in items)),
            'catalog_longest_description': max(len(i['description']) for i in items),
            'uc_details': ucs, 'fr_uc_sources': fr_sources,
            'uc_text_references_absent_from_sample': unresolved,
            'allocation_catalog_mapping_candidates': mapping,
            'documents': {name: {k: v for k, v in doc.items() if k not in ('plain_text', 'tables')} |
                         {'body_tables': len(doc['tables'])} for name, doc in
                         [('PRD文档.docx', prd_doc), ('UC.docx', uc_doc)]},
            'conversion_warning': 'Existing converter reads text but does not resolve Word numbering.xml; automatic list numbers may be lost. No rendered/page-layout verification.',
            'checks': checks,
        }
        checks['source_files_unchanged'] = all(digest(source / name) == sha for name, sha in before.items())
        if not all(checks.values()):
            raise ValueError({k: v for k, v in checks.items() if not v})
        candidate = {'status': 'pending_project_and_mapping_confirmation_do_not_import',
                     'target_project_id': None, 'design_version_id': None,
                     'source_file_sha256': before['系统划分.xlsx'],
                     'missing_import_values': ['target_project_id', 'real_import_user', 'generated_stable_ids_and_codes', 'is_allocatable_policy'],
                     'items': items, 'allocation_mapping_candidates': mapping}
        output.mkdir(parents=True, exist_ok=True)
        for name, data in [('compatibility-audit.json', summary), ('subsystem-catalog.candidate.json', candidate)]:
            (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(json.dumps({'counts': summary['counts'], 'checks': checks, 'uc_unresolved': unresolved,
                          'domain_counts': summary['domain_counts']}, ensure_ascii=False, indent=2))
    finally:
        for book in workbooks.values():
            book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    parser.add_argument('--output', type=Path, default=BASE / 'seed-candidates/xray-compatibility-20260921')
    args = parser.parse_args()
    if args.output.resolve().is_relative_to(args.source.resolve()):
        parser.error('Output must not be inside the source data directory')
    audit(args.source, args.output)
