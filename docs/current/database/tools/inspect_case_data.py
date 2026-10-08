from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json, hashlib
import openpyxl
SRC=Path(r'C:\Users\zhouguangjian\Desktop\X光\数据')
OUT=Path(__file__).resolve().parent.parent / 'seed-candidates' / 'dvr-20260921'
OUT.mkdir(parents=True,exist_ok=True)
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+NS['w']+'}'
def txt(el):
    result=[]
    for n in el.iter():
        if n.tag==W+'t': result.append(n.text or '')
        elif n.tag in (W+'br',W+'cr'): result.append('\n')
        elif n.tag==W+'tab': result.append('\t')
    return ''.join(result)
report={}
for p in SRC.iterdir():
    if not p.is_file():continue
    d={'file':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size}
    if p.suffix.lower()=='.xlsx':
        wb=openpyxl.load_workbook(p,data_only=False,read_only=False)
        d['sheets']=[]
        for ws in wb:
            rows=[]
            for row in ws:
                vals=[{'cell':c.coordinate,'value':c.value,'type':c.data_type} for c in row if c.value is not None]
                if vals: rows.append({'row':row[0].row,'cells':vals})
            d['sheets'].append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'merges':[str(r) for r in ws.merged_cells.ranges],'rows':rows})
        wb.close()
    elif p.suffix.lower()=='.docx':
        with ZipFile(p) as z:
            root=ET.fromstring(z.read('word/document.xml'))
            body=root.find('w:body',NS)
            blocks=[]; ti=0
            for child in body:
                if child.tag==W+'p':
                    text=txt(child)
                    if text: blocks.append({'type':'paragraph','text':text})
                elif child.tag==W+'tbl':
                    ti+=1; rows=[]
                    for ri,tr in enumerate(child.findall('w:tr',NS),1):
                        cells=[]
                        for ci,tc in enumerate(tr.findall('w:tc',NS),1):
                            gs=tc.find('w:tcPr/w:gridSpan',NS); vm=tc.find('w:tcPr/w:vMerge',NS)
                            cells.append({'xml_cell':ci,'text':'\n'.join(txt(x) for x in tc.findall('w:p',NS)), 'span':int(gs.get(W+'val','1')) if gs is not None else 1,'vmerge':vm.get(W+'val','continue') if vm is not None else None})
                        rows.append({'row':ri,'cells':cells})
                    blocks.append({'type':'table','table':ti,'rows':rows})
            d['blocks']=blocks
            d['tracked_insertions']=len(root.findall('.//w:ins',NS));d['tracked_deletions']=len(root.findall('.//w:del',NS))
            d['media_files']=[n for n in z.namelist() if n.startswith('word/media/')]
    report[p.name]=d

# Read-only source inspection. Outputs are staging evidence, NEVER executable SQL.
import re
from collections import Counter

def sheet(file,name):
    return next(s for s in report[file]['sheets'] if s['name'].strip()==name)
def records(file,name):
    return [{'source_row':row['row'],**{re.sub(r'\d','',c['cell']):c['value'] for c in row['cells']}} for row in sheet(file,name)['rows'][1:]]
def doc_grid(table):
    # gridSpan expands physical grid columns; vMerge propagates only explicit merges.
    grid=[]; anchors={}
    for row in table['rows']:
        result=[]; col=0
        for c in row['cells']:
            text=c['text']
            if c['vmerge']=='continue':
                if col not in anchors: raise ValueError('unmatched vertical merge')
                text=anchors[col]
            for offset in range(c['span']):
                result.append(text)
                if c['vmerge'] in ('restart','continue'):anchors[col+offset]=text
                else:anchors.pop(col+offset,None)
            col+=c['span']
        grid.append(result)
    return grid

field_map={'用例名称':'name','用例描述':'description','用例编号':'number','参与者':'actors','前置条件':'preconditions','后置条件':'postconditions','扩展用例':'extensions','HMI 交互逻辑':'hmi','法规需求':'regulations','性能需求':'performance','非功能需求':'nonfunctional'}
exception_names=['车辆状态异常','机械故障','环境突变','用户误操作（滥用）','供电异常']
exception_keys=['vehicle_state','mechanical','environmental','misuse','power']
ucs=[]
for t in [b for b in report['UC.docx']['blocks'] if b['type']=='table']:
    c={};branches=[];exceptions={};loc={}
    for row_no,row in enumerate(doc_grid(t),1):
        label=row[0].strip().rstrip('：:')
        if label in field_map:
            k=field_map[label];c[k]=row[-1];loc[k]={'table':t['table'],'row':row_no}
        elif label=='场景':
            if row[1].startswith('主事件流'):c['mainFlow']=row[-1];loc['mainFlow']={'table':t['table'],'row':row_no}
            elif row[1].startswith('分支事件流'):
                branches.append({'category':'a'+re.search(r'\d+',row[1]).group(),'body':row[-1],'source_row':row_no})
            elif row[2] in exception_names:
                exceptions[row[2]]={'body':row[-1],'source_row':row_no}
    if set(c)!=set(field_map.values())|{'mainFlow'}:raise ValueError(('missing UC fields',t['table'],set(c)))
    if set(exceptions)!=set(exception_names):raise ValueError('missing exception categories')
    c['branches']=[x['body'] for x in branches]+['']*(4-len(branches))
    c['exceptions']=[exceptions[n]['body'] for n in exception_names]
    ucs.append({'source_code':c['number'],'source_file':'UC.docx','source_table':t['table'],'content':c,'field_locations':loc,'original_branches':branches,'original_exceptions':[{'category':k,**exceptions[n]} for k,n in zip(exception_keys,exception_names)],'missing_branch_slots':list(range(len(branches)+1,5)),'projection_policy':'empty slots are adapter-only placeholders, NOT source N/A or accepted content','review_status':'pending','prd_relation_status':'proposed_import_association_not_proven_generation'})
fr_rows=records('FR列表.xlsx','行车记录仪Fr_List');duplicate_fr=records('FR划分.xlsx','行车记录仪Fr_List')
alloc_rows=records('FR划分.xlsx','行车记录仪功能需求分配表');duplicate_alloc=records('FR列表.xlsx','行车记录仪功能需求分配表')
frs=[];allocs=[];name_differences=[];uc_codes={u['source_code'] for u in ucs}
for row in fr_rows:
    sources=re.findall(r'UC-\d+-\d+-\d+',row['B'])
    if not sources or any(s not in uc_codes for s in sources):raise ValueError('unresolved UC source')
    frs.append({'source_code':row['A'],'description':row['C'],'source_text':row['B'],'uc_source_codes':sources,'source_file':'FR列表.xlsx','source_sheet':'行车记录仪Fr_List','source_row':row['source_row'],'review_status':'pending'})
for row in alloc_rows:
    allocs.append({'fr_source_code':row['C'],'subsystem_source_name':row['B'],'function_source_name':row['A'],'source_file':'FR划分.xlsx','source_sheet':'行车记录仪功能需求分配表 ','source_row':row['source_row'],'review_status':'pending'})
for row in duplicate_fr:
    mentions=re.findall(r'【(.*?)】',row['C'])
    if mentions and row['D'] not in mentions:name_differences.append({'fr_code':row['A'],'description_mentions':mentions,'allocation_name':row['D'],'resolution':'preserve both; NOT automatically aliases'})

def triple(rows,keys):return [tuple(r.get(k) for k in keys) for r in rows]
checks={'fr_copies_equal':triple(fr_rows,'ABC')==triple(duplicate_fr,'ABC'),'allocation_copies_equal':triple(alloc_rows,'ABC')==triple(duplicate_alloc,'ABC'),'allocation_D_matches_table':{r['A']:r['D'] for r in duplicate_fr}=={r['C']:r['B'] for r in alloc_rows},'unique_fr_codes':len({f['source_code'] for f in frs})==len(frs),'unique_uc_codes':len(uc_codes)==len(ucs),'allocation_fr_set_matches':{a['fr_source_code'] for a in allocs}=={f['source_code'] for f in frs},'all_uc_sources_resolve':all(s in uc_codes for f in frs for s in f['uc_source_codes']),'no_formula_cells':not any(c['type']=='f' for d in report.values() for s in d.get('sheets',[]) for row in s['rows'] for c in row['cells'])}
if not all(checks.values()):raise ValueError(checks)
# PRD body-only draft: preserve body block order and table text, not Word layout.
md=[]
def escape_cell(t):return t.replace('\\','\\\\').replace('|','\\|').replace('\n','<br>')
for b in report['PRD文档.docx']['blocks']:
    if b['type']=='paragraph':md.append(b['text'])
    else:
        rows=doc_grid(b)
        if not rows:continue
        width=max(map(len,rows));rows=[row+['']*(width-len(row)) for row in rows]
        md.append('\n'.join(['| '+' | '.join(escape_cell(v) for v in rows[0])+' |','| '+' | '.join(['---']*width)+' |']+['| '+' | '.join(escape_cell(v) for v in row)+' |' for row in rows[1:]]))
prd='\n\n'.join(md)+'\n'
(OUT/'prd-body.draft.md').write_text(prd,encoding='utf-8')
manifest=[]
for name,d in report.items():
    manifest.append({'name':name,**{k:d[k] for k in ('file','sha256','bytes')},'sheets':[{'name':s['name'],'max_row':s['max_row'],'max_column':s['max_column'],'merged_ranges':s['merges']} for s in d.get('sheets',[])],'docx_body_tables':sum(b['type']=='table' for b in d.get('blocks',[])),'tracked_insertions':d.get('tracked_insertions',0),'tracked_deletions':d.get('tracked_deletions',0)})
summary={'status':'staging_only_not_database_seed','inspected_on':'2026-09-21','sources':manifest,'counts':{'prd_documents':1,'uc':len(ucs),'fr':len(frs),'fr_uc_links':sum(len(f['uc_source_codes']) for f in frs),'allocation_objects':len({a['fr_source_code'] for a in allocs}),'allocation_targets':len(allocs),'subsystem_names':len({a['subsystem_source_name'] for a in allocs})},'checks':checks,'name_differences':name_differences,'branch_counts':[{'uc':u['source_code'],'count':len(u['original_branches']),'missing_slots':u['missing_branch_slots']} for u in ucs],'uncovered_cases':['one FR multiple subsystem targets','history changes and stale source','cross-project denial','concurrent edits','archive referenced object','retry import'],'exclusions':['封面','变更记录','01_Fr_List模板、示例','02_需求分配表'],'exclusion_reason':'FCW template rows are a different case; not part of the DVR seed'}
candidate={'format_version':1,'status':'pending_confirmation_do_not_import','target_project_id':None,'target_function_id':None,'created_via_proposal':'import','source_manifest':manifest,'required_confirmations':['isolated target project/function and real user identity','missing UC branch adapter policy','three subsystem name differences','PRD conversion and import association confirmation'],'prd':{'markdown_file':'prd-body.draft.md','source_file':'PRD文档.docx','review_status':'pending','conversion_limits':['body text/tables only','no original Word page layout, styles or automatic numbering guarantees','header logos and footer not in business body']},'ucs':ucs,'frs':frs,'allocations':allocs,'subsystems':[{'source_name':n,'code':None,'description':None,'responsibility_in':None,'responsibility_out':None,'source':'allocation target name, not a complete architecture specification'} for n in dict.fromkeys(a['subsystem_source_name'] for a in allocs)]}
(OUT/'audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'candidate-data.json').write_text(json.dumps(candidate,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'counts':summary['counts'],'checks':checks,'outputs':['audit.json','candidate-data.json','prd-body.draft.md']},ensure_ascii=False,indent=2))