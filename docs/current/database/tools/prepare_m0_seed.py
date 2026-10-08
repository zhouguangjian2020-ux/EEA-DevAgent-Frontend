"""Read supplied Word/Excel, resolve supported Word numbering, build reviewed import JSON.
Never writes source files or contacts the database. Unknown numbering fails closed.
"""
import argparse
import hashlib
import json
import re
import uuid
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+NS['w']+'}'
BASE=Path(__file__).resolve().parents[1]
CANONICAL={
 '低压电源子系统':'低压电源管理系统',
 '感知融合与定位系统':'感知融合和定位系统',
 '挡位系统':'挡位控制系统',
 'X-call系统':'X-Call系统',
 '外部连接子系统':'外部设备链接系统',
 '外部设备连接系统':'外部设备链接系统',
}
NAMESPACE=uuid.UUID('a3c6b074-06ef-4d5f-9641-1fc7c0a44024')
def uid(key):return str(uuid.uuid5(NAMESPACE,key))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def val(el,path,default=None):
 n=el.find(path,NS) if el is not None else None
 return n.get(W+'val',default) if n is not None else default

def raw_text(el):
 return ''.join(n.text or '' if n.tag==W+'t' else '\n' if n.tag in (W+'br',W+'cr') else '\t'
                for n in el.iter() if n.tag in (W+'t',W+'br',W+'cr',W+'tab'))
def alphabet(n):
 s=''
 while n:n,k=divmod(n-1,26);s=chr(97+k)+s
 return s

def roman(n):
 if not 0<n<4000:raise ValueError('Unsupported Roman number')
 s=''
 for v,c in [(1000,'m'),(900,'cm'),(500,'d'),(400,'cd'),(100,'c'),(90,'xc'),(50,'l'),(40,'xl'),(10,'x'),(9,'ix'),(5,'v'),(4,'iv'),(1,'i')]:
  while n>=v:n-=v;s+=c
 return s

class Numbering:
 def __init__(self,package):
  self.root=ET.fromstring(package.read('word/numbering.xml')) if 'word/numbering.xml' in package.namelist() else ET.Element('empty')
  self.nums={n.get(W+'numId'):n for n in self.root.findall('w:num',NS)}
  self.abstract={n.get(W+'abstractNumId'):n for n in self.root.findall('w:abstractNum',NS)}
  self.styles={}
  if 'word/styles.xml' in package.namelist():
   self.styles={n.get(W+'styleId'):n for n in ET.fromstring(package.read('word/styles.xml')).findall('w:style',NS)}
  self.counts={};self.events=[]
 def properties(self,p):
  parts=[];style_id=val(p,'w:pPr/w:pStyle');seen=set()
  while style_id:
   if style_id in seen:raise ValueError('Style inheritance loop')
   seen.add(style_id);style=self.styles.get(style_id)
   if style is None:break
   parts.insert(0,style.find('w:pPr/w:numPr',NS));style_id=val(style,'w:basedOn')
  parts.append(p.find('w:pPr/w:numPr',NS))
  result={}
  for part in parts:
   if part is not None:
    for key in ('numId','ilvl'):
     v=val(part,'w:'+key)
     if v is not None:result[key]=v
  return result
 def level(self,numid,level):
  num=self.nums[numid];abstract=self.abstract[val(num,'w:abstractNumId')]
  if abstract.find('w:numStyleLink',NS) is not None:raise ValueError('numStyleLink requires explicit implementation')
  override=num.find(f'w:lvlOverride[@w:ilvl="{level}"]',NS)
  lv=override.find('w:lvl',NS) if override is not None else None
  if lv is None:lv=abstract.find(f'w:lvl[@w:ilvl="{level}"]',NS)
  if lv is None and level==0:
   single=abstract.findall('w:lvl',NS)
   if len(single)==1 and single[0].get(W+'ilvl') is None:lv=single[0]  # Supplied PRD exporter omits ilvl on its single-level definitions.
  if lv is None:raise ValueError(('Missing numbering level',numid,level))
  start=int(val(override,'w:startOverride',val(lv,'w:start','1')))
  return lv,start
 def render(self,p,index):
  raw=raw_text(p);props=self.properties(p);numid=props.get('numId')
  if not numid or numid=='0':return raw
  level=int(props.get('ilvl','0'));lv,start=self.level(numid,level)
  # Current supplied documents use level zero only; refuse unknown multilevel semantics.
  if level!=0:raise ValueError(f'Unreviewed multilevel numbering at paragraph {index}; stop rather than guess')
  for unsupported in ['lvlRestart','isLgl','legacy']:
   if lv.find('w:'+unsupported,NS) is not None:raise ValueError(f'Unsupported active numbering feature: {unsupported}')
  key=(numid,level);self.counts[key]=self.counts.get(key,start-1)+1
  fmt=val(lv,'w:numFmt','decimal');pattern=val(lv,'w:lvlText','%1.');n=self.counts[key]
  if fmt=='decimal':mark=str(n)
  elif fmt in ('lowerLetter','upperLetter'):mark=alphabet(n);mark=mark.upper() if fmt=='upperLetter' else mark
  elif fmt in ('lowerRoman','upperRoman'):mark=roman(n);mark=mark.upper() if fmt=='upperRoman' else mark
  elif fmt=='bullet':mark=pattern
  else:raise ValueError(('Unsupported numbering format',fmt))
  if re.search(r'%[2-9]',pattern):raise ValueError('Unreviewed multilevel placeholder')
  prefix=pattern.replace('%1',mark) if fmt!='bullet' else mark
  suffix=val(lv,'w:suff','tab');sep={'tab':'\t','space':' ','nothing':''}.get(suffix)
  if sep is None:raise ValueError('Unknown number suffix')
  rendered=prefix+sep+raw
  self.events.append({'paragraph':index,'num_id':numid,'level':level,'number':n,'format':fmt,'prefix':prefix,'raw':raw,'rendered':rendered})
  return rendered

def read_word(path):
 with ZipFile(path) as z:
  root=ET.fromstring(z.read('word/document.xml'));body=root.find('w:body',NS)
  for unsupported in ['ins','del','drawing','pict','fldChar','txbxContent']:
   if body.find('.//w:'+unsupported,NS) is not None:raise ValueError(f'{path.name}: unsupported body markup {unsupported}')
  numbering=Numbering(z)
  paragraphs={id(p):numbering.render(p,i) for i,p in enumerate(body.iter(W+'p'),1)}
  blocks=[]
  for node in body:
   if node.tag==W+'p':blocks.append({'kind':'paragraph','text':paragraphs[id(node)]})
   if node.tag==W+'tbl':
    rows=[];anchors={}
    for tr in node.findall('w:tr',NS):
     row=[];col=0
     for tc in tr.findall('w:tc',NS):
      if tc.find('w:tbl',NS) is not None:raise ValueError('Nested tables require reviewed converter')
      span=int(val(tc,'w:tcPr/w:gridSpan','1'));merge=tc.find('w:tcPr/w:vMerge',NS)
      mode=merge.get(W+'val','continue') if merge is not None else None
      text='\n'.join(paragraphs[id(p)] for p in tc.findall('w:p',NS))
      if mode=='continue':
       if col not in anchors:raise ValueError('Invalid vertical merge')
       text=anchors[col]
      for offset in range(span):
       row.append(text)
       if mode in ('restart','continue'):anchors[col+offset]=text
       else:anchors.pop(col+offset,None)
      col+=span
     rows.append(row)
    blocks.append({'kind':'table','rows':rows})
  return blocks,numbering.events

def to_markdown(blocks):
 result=[]
 def cell(t):return t.replace('\\','\\\\').replace('|','\\|').replace('\n','<br>')
 for b in blocks:
  if b['kind']=='paragraph':result.append(b['text']);continue
  rows=b['rows'];width=max(map(len,rows));rows=[r+['']*(width-len(r)) for r in rows]
  result.append('\n'.join(['| '+' | '.join(cell(v) for v in rows[0])+' |','| '+' | '.join(['---']*width)+' |']+['| '+' | '.join(cell(v) for v in r)+' |' for r in rows[1:]]))
 return '\n\n'.join(result)+'\n'

FIELDS={'用例名称':'name','用例描述':'description','用例编号':'code','参与者':'actors_text','前置条件':'preconditions','后置条件':'postconditions','扩展用例':'extensions_text','HMI 交互逻辑':'hmi_logic','法规需求':'regulations','性能需求':'performance','非功能需求':'nonfunctional'}
EXCEPTIONS={'车辆状态异常':'vehicle_state','机械故障':'mechanical','环境突变':'environmental','用户误操作（滥用）':'misuse','供电异常':'power'}

def prepare(source,output,project_code,project_name):
 files=['PRD文档.docx','UC.docx','FR列表.xlsx','FR划分.xlsx','系统划分.xlsx']
 manifest={f:sha(source/f) for f in files}
 docs={f:read_word(source/f) for f in files if f.endswith('.docx')}
 ucs=[]
 for ti,b in enumerate([b for b in docs['UC.docx'][0] if b['kind']=='table'],1):
  u={'scenarios':[],'source_table':ti}
  for row in b['rows']:
   label=row[0].strip().rstrip('：:')
   if label in FIELDS:u[FIELDS[label]]=row[-1]
   elif label=='场景':
    if row[1].startswith('主事件流'):u['main_flow']=row[-1]
    elif row[1].startswith('分支事件流'):u['scenarios'].append({'kind':'branch','category':'a'+re.search(r'\d+',row[1]).group(),'body':row[-1]})
    elif row[2] in EXCEPTIONS:u['scenarios'].append({'kind':'exception','category':EXCEPTIONS[row[2]],'body':row[-1]})
   else:raise ValueError(('Unknown UC label',label))
  if set(u)-{'scenarios','source_table'} != set(FIELDS.values())|{'main_flow'}:raise ValueError('Missing UC fields')
  ucs.append(u)
 codes={u['code'] for u in ucs}
 mentions=set(re.findall(r'UC-\d+-\d+-\d+','\n'.join(v for u in ucs for v in u.values() if isinstance(v,str))))
 if mentions-codes:raise ValueError(('Still unresolved UC references',sorted(mentions-codes)))
 books={f:openpyxl.load_workbook(source/f,data_only=False) for f in files if f.endswith('.xlsx')}
 try:
  if any(c.data_type=='f' for book in books.values() for s in book for row in s for c in row):raise ValueError('Formula source requires separate review')
  def rows(file,sheet,n):
   ws=books[file][sheet]
   return [(r,tuple(ws.cell(r,c).value for c in range(1,n+1))) for r in range(2,ws.max_row+1) if ws.cell(r,1).value is not None]
  fr=rows('FR列表.xlsx','行车记录仪Fr_List',3);fr2=rows('FR划分.xlsx','行车记录仪Fr_List',4)
  a=rows('FR划分.xlsx','行车记录仪功能需求分配表 ',3);a2=rows('FR列表.xlsx','行车记录仪功能需求分配表 ',3)
  assert [v for r,v in fr]==[v[:3] for r,v in fr2] and a==a2
  assert {v[0]:v[3] for r,v in fr2}=={v[2]:v[1] for r,v in a}
  catalog=[];ws=books['系统划分.xlsx']['系统功能需求']
  for r in range(3,ws.max_row+1):
   name=ws.cell(r,3).value
   if name is None:continue
   domain=ws.cell(r,2).value
   if domain is None:
    for merged in ws.merged_cells.ranges:
     if merged.min_col<=2<=merged.max_col and merged.min_row<=r<=merged.max_row:domain=ws.cell(merged.min_row,merged.min_col).value
   desc=ws.cell(r,4).value
   if not domain or not desc:raise ValueError('Missing source catalog field')
   catalog.append({'source_row':r,'source_name':name,'name':CANONICAL.get(name,name),'description':desc,'domain_group':domain})
 finally:
  for book in books.values():book.close()
 assert len(catalog)==43 and len(ucs)==3 and len(fr)==14 and len(a)==14
 assert len({c['name'] for c in catalog})==43
 cat={c['name']:c for c in catalog}
 for _,v in a:assert CANONICAL.get(v[1],v[1]) in cat
 source_key=hashlib.sha256(json.dumps({'sources':manifest,'mapping':CANONICAL,'converter':'numbering-v1','project_code':project_code,'project_name':project_name},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
 pid=uid('project:'+project_code);fid=uid('function:'+project_code+':DVR');scheme=uid('scheme:'+project_code+':1');actor='__IMPORT_USER__'
 data={}
 def add(table,**row):data.setdefault(table,[]).append(row);return row
 def identity(typ,code):return uid(project_code+':'+typ+':'+code)
 def revision(typ,code):return uid(project_code+':'+typ+':'+code+':r1')
 def meta(file,**where):return {'file':file,'sha256':manifest[file],**where,'seed_checksum':source_key,'converter':'numbering-v1'}
 def obj(typ,code,file,where):
  aid=identity(typ,code);rid=revision(typ,code)
  add('artifacts',id=aid,project_id=pid,function_id=fid,type_code=typ,code=code,head_revision_id=rid,created_by=actor,metadata={'seed_checksum':source_key})
  add('artifact_revisions',id=rid,project_id=pid,function_id=fid,artifact_id=aid,type_code=typ,revision_no=1,type_version_id=uid('type:'+typ+':m0-v1'),created_via='import',created_by=actor,metadata=meta(file,**where),change_summary='用户确认的样本导入；保存待审阅，不伪造历史AI生成')
  return aid,rid
 for typ,title,key in [('PRD','PRD','prd'),('UC','用例','uc'),('FR','功能需求','fr'),('FR_ALLOCATION','FR分配','fr-allocation')]:
  add('artifact_types',code=typ,display_name=title,workspace_key=key)
  if typ=='PRD':props={'markdown_body':{'type':'string'}};req=['markdown_body']
  elif typ=='UC':
   props={v:{'type':'string'} for v in FIELDS.values() if v!='code'};props['main_flow']={'type':'string'}
   props['scenarios']={'type':'array','items':{'type':'object','properties':{'kind':{'enum':['branch','exception']},'category':{'type':'string'},'body':{'type':'string'}},'required':['kind','category','body'],'additionalProperties':False}}
   req=list(props)
  elif typ=='FR':props={'description':{'type':'string'},'uc_sources':{'type':'array','minItems':1,'items':{'type':'string','format':'uuid'}}};req=list(props)
  else:props={'subsystem_ids':{'type':'array','uniqueItems':True,'items':{'type':'string','format':'uuid'}}};req=list(props)
  # Registration is a draft transport schema; separate review rules retain missing branches.
  add('artifact_type_versions',id=uid('type:'+typ+':m0-v1'),type_code=typ,version_no=1,created_by=actor,content_schema={'$schema':'https://json-schema.org/draft/2020-12/schema','type':'object','properties':props,'required':req,'additionalProperties':False},field_catalog=[{'key':k,'path':k,'type':v.get('type','string')} for k,v in props.items()])
 add('projects',id=pid,code=project_code,name=project_name,description='数据库先行样本项目；尚未接前后端API',metadata={'seed_checksum':source_key,'source_manifest':manifest,'canonical_mapping':CANONICAL},created_by=actor)
 add('business_functions',id=fid,project_id=pid,code='DVR-301',name='行车记录仪DVR功能',metadata={'source_allocation_function_name':'行车录制功能','association':'controlled_import'},created_by=actor)
 add('project_members',project_id=pid,user_id=actor,added_by=actor)
 role=uid('role:m0-seed-reader')
 add('access_roles',id=role,code='m0_seed_reader',name='样本只读角色',description='本阶段只登记只读能力，不等于已实现应用鉴权')
 for resource in ['project','function','PRD','UC','FR','FR_ALLOCATION','subsystem_design']:
  code=resource.lower()+'.read';add('permission_definitions',code=code,display_name=resource+'读取',resource_kind=resource,action='read',allowed_scopes=['project'])
  add('role_permissions',role_id=role,permission_code=code)
 add('access_bindings',id=uid(project_code+':importer-binding'),project_id=pid,user_id=actor,scope_kind='project',role_id=role,created_by=actor,updated_by=actor)
 add('project_subsystem_designs',project_id=pid,name='项目预置子系统划分',current_version_id=scheme,created_by=actor,updated_by=actor,metadata={'source_manifest':manifest,'canonical_mapping':CANONICAL})
 add('subsystem_design_versions',id=scheme,project_id=pid,version_no=1,created_via='import',created_by=actor,change_summary='采用用户确认名称预置43项；不是AI自动划分',metadata=meta('系统划分.xlsx'))
 for i,c in enumerate(catalog,1):
  sid=uid(project_code+':subsystem:'+str(c['source_row']));c['id']=sid
  add('subsystems',id=sid,project_id=pid,code=f'SS-{i:03}',created_by=actor)
  add('subsystem_version_items',project_id=pid,design_version_id=scheme,subsystem_id=sid,name=c['name'],description=c['description'],domain_group=c['domain_group'],sort_order=i-1,is_allocatable=True,source_metadata=meta('系统划分.xlsx',sheet='系统功能需求',row=c['source_row'],original_name=c['source_name'],name_rule='user_confirmed_20260921',allocatable_rule='initial_sample_all_43_selectable'))
 prdid,prdrid=obj('PRD','PRD-DVR-301','PRD文档.docx',{})
 add('prds',artifact_id=prdid,project_id=pid,function_id=fid)
 prdtext=to_markdown(docs['PRD文档.docx'][0])
 add('prd_versions',revision_id=prdrid,project_id=pid,function_id=fid,prd_id=prdid,markdown_body=prdtext)
 for u in ucs:
  aid,rid=obj('UC',u['code'],'UC.docx',{'table':u['source_table'],'prd_relation':'import_association_not_historical_AI','missing_branch_slots':[f'a{i}' for i in range(1,5) if f'a{i}' not in {s['category'] for s in u['scenarios'] if s['kind']=='branch'}]})
  add('use_cases',artifact_id=aid,project_id=pid,function_id=fid,prd_id=prdid)
  add('uc_versions',revision_id=rid,project_id=pid,function_id=fid,uc_id=aid,prd_id=prdid,prd_revision_id=prdrid,**{k:v for k,v in u.items() if k not in ['scenarios','source_table','code']})
  counts=Counter()
  for s in u['scenarios']:
   add('uc_scenario_items',project_id=pid,function_id=fid,uc_id=aid,uc_revision_id=rid,item_id=uid(aid+':'+s['kind']+':'+s['category']),position=counts[s['kind']],**s);counts[s['kind']]+=1
 for r,v in fr:
  code,source_text,description=v;aid,rid=obj('FR',code,'FR列表.xlsx',{'sheet':'行车记录仪Fr_List','row':r,'original_source_text':source_text})
  add('functional_requirements',artifact_id=aid,project_id=pid,function_id=fid)
  add('fr_versions',revision_id=rid,project_id=pid,function_id=fid,fr_id=aid,description=description)
  for uc in re.findall(r'UC-\d+-\d+-\d+',source_text):
   assert uc in codes
   add('fr_uc_sources',project_id=pid,function_id=fid,fr_revision_id=rid,uc_id=identity('UC',uc),uc_revision_id=revision('UC',uc),rationale='按源Excel编号关联本批导入UC版本，不虚构原历史版本')
 for r,v in a:
  function_name,target,frcode=v;canonical=CANONICAL.get(target,target);code='ALLOC-'+frcode
  aid,rid=obj('FR_ALLOCATION',code,'FR划分.xlsx',{'sheet':'行车记录仪功能需求分配表 ','row':r,'original_target_name':target,'canonical_target_name':canonical})
  add('fr_allocations',artifact_id=aid,project_id=pid,function_id=fid,fr_id=identity('FR',frcode))
  add('allocation_versions',revision_id=rid,project_id=pid,function_id=fid,allocation_id=aid,fr_id=identity('FR',frcode),fr_revision_id=revision('FR',frcode),design_version_id=scheme,note='用户确认名称映射；原分配目标：'+target)
  add('allocation_targets',project_id=pid,function_id=fid,allocation_revision_id=rid,design_version_id=scheme,subsystem_id=cat[canonical]['id'],rationale='按源分配表及用户确认名称导入')
 add('security_audit_events',id=uid(project_code+':seed-audit:'+source_key),project_id=pid,actor_id=actor,event_type='sample.import',scope_kind='project',after_state={'source_checksum':source_key,'counts':{k:len(v) for k,v in data.items()}},reason='操作人执行受控样本导入；原始文件不变',request_id='xray:'+source_key)
 assert len(data['fr_uc_sources'])==16
 assert all(sha(source/f)==h for f,h in manifest.items())
 output.mkdir(parents=True,exist_ok=True)
 result={'format_version':1,'project_code':project_code,'project_name':project_name,'project_id':pid,'function_id':fid,'source_checksum':source_key,'source_manifest':manifest,'confirmed_mapping':CANONICAL,'data':data}
 (output/'seed-data.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 (output/'prd-numbered.preview.md').write_text(prdtext,encoding='utf-8')
 report={'source_manifest':manifest,'source_files_unchanged':True,'unresolved_uc_references':sorted(mentions-codes),'numbering':{f:events for f,(_,events) in docs.items()},'counts':{k:len(v) for k,v in data.items()},'conversion_scope':'body paragraphs and tables, automatic numbering resolved; Word layout/attachments not persisted','generated_codes':'SS-001..SS-043 generated once for this project, not source identifiers','default_policy':'all 43 imported catalog entries allocatable; no content accepted'}
 (output/'preparation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(json.dumps({'unresolved_uc_references':sorted(mentions-codes),'numbered_paragraphs':{f:len(events) for f,(_,events) in docs.items()},'counts':report['counts']},ensure_ascii=False,indent=2))

if __name__=='__main__':
 p=argparse.ArgumentParser(description=__doc__)
 p.add_argument('--source',type=Path,default=Path(r'C:\Users\zhouguangjian\Desktop\X光\数据'))
 p.add_argument('--output',type=Path,default=BASE/'m0-20260921/data')
 p.add_argument('--project-code',default='XRAY-DVR-SAMPLE')
 p.add_argument('--project-name',default='X光DVR样本项目')
 a=p.parse_args()
 if a.output.resolve().is_relative_to(a.source.resolve()):p.error('Output must be outside source directory')
 if not re.fullmatch(r'[A-Z0-9][A-Z0-9_-]{2,63}',a.project_code):p.error('Invalid project code')
 prepare(a.source,a.output,a.project_code,a.project_name)

